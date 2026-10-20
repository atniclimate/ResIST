"""Real-document helpers for tests (D-052, D-058; design 4.4). No tests live here.

Inputs are only the four committed fixtures in tests/fixtures/atni/, the pool through mode=ro, and files derived
from them at test time. Nothing here prints or asserts on document text: messages carry file names, counts and
SHA-256 values only.
"""

import hashlib
import os
import re
import sqlite3
from functools import lru_cache
from pathlib import Path

import pytest

from resist.errors import QuarantinedError, TierRefused
from resist.models import DocRef
from resist.services import policy

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = ROOT / "tests" / "fixtures" / "atni"
README = FIXTURES / "README.md"
# Captured at import, before any fixture points RESIST_DATA_DIR at a scratch pool.
REAL_POOL = Path(os.environ.get("RESIST_DATA_DIR") or ROOT / "datapool")
# The exact pattern from scripts/build_pages.py:25.
CONTACT = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+|\(?\b\d{3}\)?[-. ]\d{3}[-. ]\d{4}\b")

_ROW = re.compile(r"^\|\s*`([^`]+\.pdf)`\s*\|\s*`([0-9a-f]{64})`\s*\|\s*`([^`]+)`", re.M)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(Path(path).read_bytes())


@lru_cache(maxsize=1)
def readme_table() -> tuple[tuple[str, str, str], ...]:
    """(file, sha256, source url) for every row of the fixture README table."""
    return tuple(_ROW.findall(README.read_text(encoding="utf-8")))


def readme_rows() -> dict[str, str]:
    return {f: sha for f, sha, _url in readme_table()}


@lru_cache(maxsize=1)
def _fixtures() -> tuple[tuple[str, Path], ...]:
    out = []
    for f, sha in readme_rows().items():
        p = FIXTURES / f
        if sha256_file(p) != sha:
            raise AssertionError(f"fixture bytes disagree with README row: {f}")
        out.append((sha, p))
    return tuple(out)


def fixtures() -> dict[str, Path]:
    """sha256 -> committed fixture path, every README row checked against the file bytes."""
    return dict(_fixtures())


def fixture(name: str) -> tuple[str, Path]:
    """(sha256, path) of one committed fixture by file name."""
    sha = readme_rows()[name]
    return sha, fixtures()[sha]


def truncated_copy(name: str, out_dir: Path) -> tuple[str, Path]:
    """(sha256, path) of the first half of a committed fixture, written at test time. Its SHA-256 is in no
    manifest, so with the manifests present it derives T3 (D-038)."""
    data = fixture(name)[1].read_bytes()
    out_dir.mkdir(parents=True, exist_ok=True)
    cut = out_dir / f"truncated-{name}"
    cut.write_bytes(data[: len(data) // 2])
    return sha256_file(cut), cut


def _policy(sha: str) -> policy.SourcePolicy:
    """source_policy when the manifests are present; otherwise README admission for a committed fixture (D-058)."""
    if not policy.load_manifests().present and sha in readme_rows().values():
        return policy.SourcePolicy(sha, "atni", "atni-fixture", "resolution", "T0", "public-atni", "D-058", False)
    return policy.source_policy(sha)


def doc_ref(sha: str) -> DocRef:
    return _policy(sha).doc_ref()


def original(sha: str) -> Path:
    """The original bytes' path. The tier gate runs before any byte of the document is read."""
    p = _policy(sha)
    if p.quarantined:
        raise QuarantinedError(f"quarantined: {sha} {p.decision_id}")
    if p.tier == "T3":
        raise TierRefused(f"tier-refused: {sha} {p.tier}")
    if sha in readme_rows().values():
        return fixtures()[sha]
    res_dir = policy.default_res_dir()
    for row in policy.load_manifests().atni.get(sha, ()):
        f = res_dir / (row.get("local_path") or "")
        if row.get("local_path") and f.is_file() and sha256_file(f) == sha:
            return f
    pytest.skip(f"pool original absent: {sha}")


DOCX_MT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def fixture_clauses(name: str) -> tuple[str, ...]:
    """Clause texts of a committed fixture's parse, document order (whereas then resolved); computed at test time."""
    from src.ingestion.intake import extract_text_from_native_pdf
    from src.ingestion.parser import parse_resolution

    sha, path = fixture(name)
    pr = parse_resolution(extract_text_from_native_pdf(path), sha)
    return tuple(s["text"] for s in [*pr.whereas_statements, *pr.resolved_statements])


def derived_docx(name: str, out_dir: Path, layout: str = "body") -> tuple[str, Path, tuple[str, ...]]:
    """Write the fixture's clauses into out_dir/derived-<stem>-<layout>.docx with python-docx (never committed;
    out_dir is a tmp_path or scratch-pool path). layout "body": one paragraph per clause. layout "mixed": clause 0 as
    the section header paragraph, 1 a body paragraph, 2 and 3 a 1x2 table, 4 a body paragraph, 5 the footer
    paragraph. layout "text": one paragraph per line of the fixture's whole native text (markers included, so the
    parser finds its clauses). Returns (sha256 of the saved bytes, path, the expected extracted lines in order)."""
    import docx

    d = docx.Document()
    if layout == "text":
        from src.ingestion.intake import extract_text_from_native_pdf

        clauses = tuple(extract_text_from_native_pdf(fixture(name)[1]).split("\n"))
    else:
        clauses = fixture_clauses(name)
    if layout in ("body", "text"):
        for c in clauses:
            d.add_paragraph(c)
        expected = clauses
    elif layout == "mixed":
        sec = d.sections[0]
        sec.header.paragraphs[0].text = clauses[0]
        d.add_paragraph(clauses[1])
        row = d.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = clauses[2], clauses[3]
        d.add_paragraph(clauses[4])
        sec.footer.paragraphs[0].text = clauses[5]
        expected = clauses[:6]
    else:
        raise ValueError(f"unknown layout: {layout}")
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"derived-{Path(name).stem}-{layout}.docx"
    d.save(str(out))
    return sha256_file(out), out, tuple(expected)


def docx_ref(pool: Path, sha: str) -> DocRef:
    """A derived DOCX's SHA-256 is in no manifest, so it is T3 (D-038): writes pool/pool.yaml "pool_max_tier: T3"
    and returns DocRef(sha, DOCX_MT, "atni", "T3")."""
    (pool / "pool.yaml").write_text("pool_max_tier: T3\n", encoding="utf-8")
    return DocRef(sha, DOCX_MT, "atni", "T3")


def pool_db() -> Path:
    p = REAL_POOL / "db" / "resolutions.db"
    if not p.is_file():
        pytest.skip("pool absent")
    return p


def ro(path: Path) -> sqlite3.Connection:
    """Read-only connection through a mode=ro URI; the file is never opened writable."""
    return sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro", uri=True)
