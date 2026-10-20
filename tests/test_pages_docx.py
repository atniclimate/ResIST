"""Opens the site's Word export (pages/docx.js, D-054) with python-docx and checks it reads like ATNI's template.

Node writes the documents (tests/js/docx-fixtures.mjs); this test skips cleanly when Node is missing.
"""
import json
import re
import shutil
import sqlite3
import subprocess
import zipfile
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.oxml.ns import nsmap  # noqa: E402  (includes w14)
from lxml import etree  # noqa: E402  (python-docx depends on lxml)

from resist.config import data_dir  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
NODE = shutil.which("node")
DB = data_dir() / "exchange" / "pages" / "resolutions.db"  # RESIST_DATA_DIR, read at import (D-064 item 4)
LEADS = ("WHEREAS,", "THEREFORE, BE IT RESOLVED,", "BE IT FURTHER RESOLVED,", "BE IT FINALLY RESOLVED,")

pytestmark = pytest.mark.skipif(NODE is None, reason="node is not installed")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("docx")
    run = subprocess.run([NODE, "tests/js/docx-fixtures.mjs", str(out)], cwd=ROOT, capture_output=True, text=True,
                         timeout=120)
    assert run.returncode == 0, run.stderr[-2000:]
    return out


def xp(el, path):
    """XPath with the OOXML prefixes, on python-docx and plain lxml elements alike."""
    return etree._Element.xpath(el, path, namespaces=nsmap)


def text(el):
    """Visible text of an element, including what sits inside content controls."""
    return "".join(xp(el, ".//w:t/text()"))


def body_paras(document):
    return xp(document.element.body, ".//w:p")


def placeholders(element):
    return {t for t in xp(element, ".//w:sdt[w:sdtPr/w:showingPlcHdr]/w:sdtPr/w:tag/@w:val")}


def checkboxes(element):
    return {xp(sdt, "string(w:sdtPr/w:tag/@w:val)"):
            xp(sdt, "string(w:sdtPr/w14:checkbox/w14:checked/@w14:val)")
            for sdt in xp(element, ".//w:sdt[w:sdtPr/w14:checkbox]")}


def para_with(document, start):
    for p in body_paras(document):
        if text(p).startswith(start):
            return p
    raise AssertionError(f"no paragraph starting {start!r}")


def bold(r):
    return bool(xp(r, "w:rPr/w:b"))


def test_every_part_parses_and_is_clean(built):
    for name in ("empty.docx", "full.docx"):
        with zipfile.ZipFile(built / name) as z:
            assert z.testzip() is None
            for part in z.namelist():
                data = z.read(part)
                if part.endswith((".xml", ".rels")):
                    etree.fromstring(data)
                    s = data.decode("utf-8")
                    assert "\u2014" not in s, part
                    assert "@" not in s, part
                    assert not re.search(r"\(?\d{3}\)?[ .-]\d{3}[ .-]\d{4}", s), f"{part}: phone number"


def test_defaults_are_times_new_roman_12pt(built):
    d = docx.Document(built / "empty.docx")
    rpr = xp(d.styles.element, "w:docDefaults/w:rPrDefault/w:rPr")[0]
    assert xp(rpr, "string(w:rFonts/@w:ascii)") == "Times New Roman"
    assert xp(rpr, "string(w:sz/@w:val)") == "24"
    assert xp(rpr, "string(w:lang/@w:val)") == "en-US"
    section = d.sections[0]
    assert (section.page_width.twips, section.page_height.twips) == (12240, 15840)
    assert {section.left_margin.twips, section.right_margin.twips, section.top_margin.twips,
            section.bottom_margin.twips} == {1440}


def test_full_resolution_reads_like_the_template(built):
    d = docx.Document(built / "full.docx")
    paras = [text(p) for p in body_paras(d)]
    assert "ATNI 2026 Annual Convention" in paras
    assert "Snoqualmie, Washington" in paras
    assert "RESOLUTION # 2026 – DRFT-XX" in paras
    assert "“Supporting Responsible Governance of Youth Use of Artificial Intelligence Chatbots in Schools and " \
           "Education”" in paras
    assert "PREAMBLE" in paras
    assert "Draft 1 Due to the ATNI Workroom by August 28, 2026" in paras
    heading = para_with(d, "ATNI 2026")
    assert xp(heading, "string(w:pPr/w:jc/@w:val)") == "center"
    assert xp(heading, "string(w:r[1]/w:rPr/w:sz/@w:val)") == "32"

    clauses = [p for p in body_paras(d) if text(p).startswith(LEADS)]
    whereas = [p for p in clauses if text(p).startswith("WHEREAS,")]
    resolved = [p for p in clauses if not text(p).startswith("WHEREAS,")]
    assert len(whereas) == 4 and len(resolved) == 4
    for p in clauses:
        runs = xp(p, "w:r")
        assert text(runs[0]) in LEADS and bold(runs[0])
        assert not any(bold(r) for r in runs[1:])
        assert xp(p, "string(w:pPr/w:ind/@w:firstLine)") == "720"
    assert all(text(p).endswith("; and") for p in whereas[:-1]) and text(whereas[-1]).endswith("; now")
    assert all(text(p).endswith("; and") for p in resolved[:-1]) and text(resolved[-1]).endswith(".")
    assert text(resolved[0]).startswith("THEREFORE, BE IT RESOLVED, that ATNI urges both Tribal and Public Schools")
    assert text(resolved[-1]).startswith("BE IT FINALLY RESOLVED, that this will be the policy of ATNI")
    assert "PAGE 2" not in text(whereas[-1])

    cert = para_with(d, "CERTIFICATION")
    assert xp(cert, "w:r/w:rPr/w:u") and xp(cert, "w:r/w:rPr/w:b")
    assert ("The foregoing Resolution was adopted at the Affiliated Tribes of Northwest Indians 2026 Annual Convention "
            "held at the Snoqualmie Casino and Hotel in Snoqualmie, Washington on September 13-17, 2026, with a quorum "
            "present.") in paras

    resolution = xp(d.element.body, "./w:p[w:pPr/w:sectPr][1]/preceding-sibling::*")
    plc = set().union(*(placeholders(e) for e in resolution))
    assert plc == {"President's name", "Secretary's name"}, "known details are plain runs"


def test_empty_draft_keeps_fillable_placeholders(built):
    d = docx.Document(built / "empty.docx")
    plc = placeholders(d.element.body)
    for tag in ("Convention", "City, State", "Venue", "Convention dates", "Draft 1 due date",
                "Final draft due date and time", "Resolution title", "President's name", "Secretary's name",
                "Author", "Sponsor", "Author's email", "Sponsor's cell phone"):
        assert tag in plc, tag
    sdt = xp(d.element.body, ".//w:sdt[w:sdtPr/w:tag/@w:val='Convention dates']")[0]
    assert xp(sdt, "w:sdtPr/w:text") and xp(sdt, "string(w:sdtPr/w:alias/@w:val)") == "Convention dates"
    assert text(sdt) == "[Convention dates]"
    assert xp(sdt, "string(.//w:rStyle/@w:val)") == "PlaceholderText"
    assert len(d.sections) == 2
    assert "this form must be completed in its entirety" not in text(d.element.body)


def test_sections_headers_and_footers(built):
    d = docx.Document(built / "full.docx")
    assert len(d.sections) == 3
    s1, s2, s3 = d.sections
    assert s1.different_first_page_header_footer
    assert text(s1.header._element) == "AFFILIATED TRIBES OF NORTHWEST INDIANS RESOLUTION #2026-DRFT-XX"
    assert text(s1.first_page_header._element) == "" and text(s1.first_page_footer._element) == ""
    footer = text(s1.footer._element)
    assert footer.startswith("2026 ANNUAL CONVENTION") and "PAGE" in footer
    assert xp(s1.footer._element, ".//w:fldSimple[contains(@w:instr, 'PAGE')]")
    assert text(s2.first_page_header._element) == ("AFFILIATED TRIBES OF NORTHWEST INDIANS"
                                                   "RESOLUTION SUBMISSION FORM")
    assert xp(s2.first_page_header._element, ".//w:drawing"), "seal in the form header"
    assert text(s2.first_page_footer._element) == ("* If different than the author. Sponsor must be an ATNI Member in "
                                                   "good standing.ATNIANN2026")
    assert text(s2.header._element) == "" and text(s2.footer._element) == ""
    assert text(s3.header._element) == "DRAFTING NOTES, NOT PART OF THE RESOLUTION"

    empty = docx.Document(built / "empty.docx").sections[0]
    assert "Convention" in placeholders(empty.footer._element), "unknown convention stays fillable in the footer"


def test_form_checkboxes(built):
    full = checkboxes(docx.Document(built / "full.docx").element.body)
    assert full["Forward to NCAI: Yes"] == "1" and full["Forward to NCAI: No"] == "0"
    assert full["Committee: Education"] == full["Committee: Youth"] == full["Committee: Other"] == "1"
    assert full["Committee: Health"] == "0" and full["Attachments: Yes"] == full["Attachments: No"] == "0"
    assert full["NCAI: Sponsor verifies information"] == "0", "the NCAI form is present when forwarding"
    assert len([k for k in full if k.startswith("Committee: ")]) == 23
    empty = checkboxes(docx.Document(built / "empty.docx").element.body)
    assert empty["Forward to NCAI: No"] == "1" and "NCAI: Sponsor verifies information" not in empty
    d = docx.Document(built / "full.docx")
    assert "Climate" in text(para_with(d, "☒ Other"))
    for sdt in xp(d.element.body, ".//w:sdt[w:sdtPr/w14:checkbox]"):
        glyph = text(sdt)
        assert glyph in ("☐", "☒")
        assert xp(sdt, "string(.//w:rFonts/@w:ascii)") == "MS Gothic"


def test_drafting_notes(built):
    d = docx.Document(built / "full.docx")
    notes = text(d.element.body)
    assert "Prepared with ResIST." in notes
    # Since LR-D (D-071) the stored clause carries no running header, so no "page header removed" note applies
    assert "Verbatim" in notes and "SHA-256 5b62a9708e4081fd" in notes
    assert "Read by Policy Sentinel on 10/19/2026; text SHA-256 b3ee0b8f7d68c588" in notes
    assert "Consultation wording: “Executive Order 13175" in notes
    assert "Drafter's note: Check the consultation section before the committee meets." in notes
    assert "are not legal conclusions." in notes
    links = {r.target_ref for r in d.part.rels.values() if r.is_external}
    assert "https://www.federalregister.gov/d/2025-23880" in links
    assert any(u.startswith("https://atnitribes.org/wp-content/uploads/2026/09/20260917_26-51_") for u in links)


@pytest.mark.skipif(not DB.exists(), reason="the site database is not built")
def test_fixture_clauses_are_verbatim_from_the_database(built):
    fixture = json.loads((built / "statements.json").read_text(encoding="utf-8"))
    con = sqlite3.connect(f"file:{DB.as_posix()}?mode=ro", uri=True)
    try:
        for kind, index, stored in fixture["statements"]:
            row = con.execute('SELECT s.text FROM statement s JOIN resolution r ON r.id = s.resolution_id '
                              'WHERE r.document_id = ? AND s.statement_type = ? AND s."index" = ?',
                              (fixture["documentId"], kind, index)).fetchone()
            assert row and row[0] == stored, (kind, index)
    finally:
        con.close()
